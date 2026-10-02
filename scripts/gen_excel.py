# -*- coding: utf-8 -*-
"""生成 1200 词英语词汇表 Excel（小升初必备）"""
import json, unicodedata, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Side
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
vocab = json.load(open(os.path.join(BASE, 'data', 'vocab-full.json'), 'r', encoding='utf-8'))
levels = json.load(open(os.path.join(BASE, 'data', 'levels.json'), 'r', encoding='utf-8'))
topic_title = {l['topic']: l['title'] for l in levels}

def dw(value):
    return sum(2 if unicodedata.east_asian_width(ch) in ('F', 'W') else 1 for ch in str(value or ''))

wb = Workbook()
ws = wb.active
ws.title = '词汇总表'

headers = ['序号', '英文', '音标', '词性', '中文释义', '主题', '例句', '例句译文']
ws.append(headers)

# 表头样式
hfill = PatternFill('solid', fgColor='17365D')
thin = Side(style='thin', color='B8C7D9')
for c in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=c)
    cell.fill = hfill
    cell.font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = __import__('openpyxl').styles.Border(bottom=thin)

# 数据
for i, v in enumerate(vocab, 1):
    topic = topic_title.get(v['topic'], v['topic'])
    ws.append([i, v['w'], v.get('p', ''), v.get('pos', ''), v.get('m', ''),
               topic, v.get('e', ''), v.get('ec', '')])

# 列宽适配
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = max(8, min(max(dw(c.value) for c in ws[col] if c.value is not None) + 3, 50))

# 对齐：序号居中，其余左对齐垂直居中
for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.font = Font(name='微软雅黑', size=10)
        if cell.column == 1:
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)

ws.freeze_panes = 'A2'
ws.sheet_view.showGridLines = True

out = os.path.join(os.path.dirname(BASE), '1200词词汇表.xlsx')
wb.save(out)
print('已生成:', out, '行数(含表头):', ws.max_row, '列数:', ws.max_column)
